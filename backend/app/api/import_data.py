"""SLAngel — CSV/JSON Import API Routes"""

import csv
import io
import json
from datetime import datetime, timedelta
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.models import (
    Application, Document, AuditLog, TimelineEvent,
    ApplicationStatus, VerificationStatus
)
from app.services.sla_engine import update_application_sla
from app.services.prediction import update_application_prediction
from app.services.priority_engine import update_application_priority
from app.services.alert_engine import generate_alerts_for_application
from app.api.applications import _generate_app_number

router = APIRouter(prefix="/api/import", tags=["Import"])


REQUIRED_CSV_COLUMNS = {
    "applicant_name", "service_type", "department"
}


def _import_single_record(record: dict, db: Session) -> tuple:
    """Import a single record. Returns (application, error_message)."""
    try:
        # Validate required fields
        if not record.get("applicant_name"):
            return None, "Missing applicant_name"
        if not record.get("service_type"):
            return None, "Missing service_type"
        if not record.get("department"):
            return None, "Missing department"

        now = datetime.utcnow()
        app_number = _generate_app_number()

        # Parse submission date
        submission_date = now
        if record.get("submission_date"):
            try:
                submission_date = datetime.strptime(record["submission_date"], "%Y-%m-%d")
            except ValueError:
                try:
                    submission_date = datetime.strptime(record["submission_date"], "%d/%m/%Y")
                except ValueError:
                    pass

        sla_days = int(record.get("sla_days", 15))

        application = Application(
            application_number=app_number,
            applicant_name=record["applicant_name"],
            applicant_contact=record.get("applicant_contact", ""),
            service_type=record["service_type"],
            department=record["department"],
            district=record.get("district", ""),
            stage=record.get("stage", "Document Verification"),
            submission_date=submission_date,
            sla_days=sla_days,
            status=record.get("status", ApplicationStatus.SUBMITTED.value),
            verification_status=record.get("verification_status", VerificationStatus.PENDING.value),
            purpose=record.get("purpose", ""),
            aadhaar_status=record.get("aadhaar_status", ""),
            annual_income=record.get("annual_income", ""),
            last_action_date=now,
            created_at=now,
            updated_at=now,
        )
        db.add(application)
        db.flush()

        # Timeline
        timeline = TimelineEvent(
            application_id=application.id,
            title="Application imported via data import",
            date_label=now.strftime("%d %b %Y"),
            created_at=now,
        )
        db.add(timeline)

        # Audit
        audit = AuditLog(
            application_id=application.id,
            action="Application imported",
            new_value=app_number,
            timestamp=now,
        )
        db.add(audit)

        db.flush()

        # Calculate
        update_application_sla(application)
        update_application_prediction(application, db)
        update_application_priority(application)
        generate_alerts_for_application(application, db)

        return application, None

    except Exception as e:
        return None, str(e)


@router.post("/csv")
async def import_csv(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """Import applications from a CSV file."""
    if not file.filename.endswith(".csv"):
        raise HTTPException(status_code=400, detail="File must be a CSV")

    content = await file.read()
    text = content.decode("utf-8")
    reader = csv.DictReader(io.StringIO(text))

    # Validate columns
    if reader.fieldnames:
        missing = REQUIRED_CSV_COLUMNS - set(reader.fieldnames)
        if missing:
            raise HTTPException(
                status_code=400,
                detail=f"Missing required columns: {', '.join(missing)}"
            )

    total = 0
    imported = 0
    failed = 0
    errors = []

    for i, row in enumerate(reader, start=1):
        total += 1
        app, error = _import_single_record(row, db)
        if error:
            failed += 1
            errors.append(f"Row {i}: {error}")
        else:
            imported += 1

    db.commit()

    return {
        "total_rows": total,
        "imported": imported,
        "failed": failed,
        "errors": errors[:20],  # Limit error messages
    }


@router.post("/json")
async def import_json(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """Import applications from a JSON file."""
    if not file.filename.endswith(".json"):
        raise HTTPException(status_code=400, detail="File must be JSON")

    content = await file.read()
    try:
        data = json.loads(content.decode("utf-8"))
    except json.JSONDecodeError:
        raise HTTPException(status_code=400, detail="Invalid JSON format")

    if not isinstance(data, list):
        data = [data]

    total = len(data)
    imported = 0
    failed = 0
    errors = []

    for i, record in enumerate(data, start=1):
        app, error = _import_single_record(record, db)
        if error:
            failed += 1
            errors.append(f"Record {i}: {error}")
        else:
            imported += 1

    db.commit()

    return {
        "total_rows": total,
        "imported": imported,
        "failed": failed,
        "errors": errors[:20],
    }


@router.post("/excel")
async def import_excel(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """Import applications from an Excel (.xlsx) file."""
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="File must be an Excel file (.xlsx or .xls)")

    try:
        import openpyxl
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="openpyxl is not installed. Excel import is not available."
        )

    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise HTTPException(status_code=400, detail="Excel file is empty or has no data rows")

    headers = [str(h).strip().lower().replace(" ", "_") if h else f"col_{i}" for i, h in enumerate(rows[0])]

    # Validate required columns
    missing = REQUIRED_CSV_COLUMNS - set(headers)
    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"Missing required columns: {', '.join(missing)}"
        )

    total = 0
    imported = 0
    failed = 0
    errors = []

    for i, row in enumerate(rows[1:], start=2):
        total += 1
        record = {}
        for j, header in enumerate(headers):
            if j < len(row):
                val = row[j]
                record[header] = str(val) if val is not None else ""
            else:
                record[header] = ""

        app, error = _import_single_record(record, db)
        if error:
            failed += 1
            errors.append(f"Row {i}: {error}")
        else:
            imported += 1

    db.commit()
    wb.close()

    return {
        "total_rows": total,
        "imported": imported,
        "failed": failed,
        "errors": errors[:20],
    }


@router.post("/pdf")
async def import_pdf(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """Import applications from a PDF file."""
    if not file.filename.endswith(".pdf"):
        raise HTTPException(status_code=400, detail="File must be a PDF file (.pdf)")

    content = await file.read()
    text = ""

    # Try pdfplumber or pypdf to extract text
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page in pdf.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
    except Exception:
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(content))
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Could not parse PDF file: {str(e)}")

    if not text.strip():
        raise HTTPException(status_code=400, detail="PDF contains no extractable text")

    # Try to parse CSV or JSON or Key-Value records inside PDF text
    records = []
    
    # Check if text contains JSON array or JSON objects
    if "[" in text and "]" in text:
        try:
            start_idx = text.find("[")
            end_idx = text.rfind("]") + 1
            json_str = text[start_idx:end_idx]
            records = json.loads(json_str)
        except Exception:
            pass

    # If not JSON, try line-by-line CSV or Key-Value parsing
    if not records:
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        
        # Check if first line is CSV header
        if any(h in lines[0].lower() for h in ["applicant", "service", "department"]):
            header = [h.strip().lower().replace(" ", "_") for h in lines[0].split(",")]
            for line in lines[1:]:
                parts = [p.strip() for p in line.split(",")]
                if len(parts) >= len(REQUIRED_CSV_COLUMNS):
                    rec = {header[i]: parts[i] for i in range(min(len(header), len(parts)))}
                    records.append(rec)
        else:
            # Parse key-value block format (e.g. Applicant Name: ..., Service Type: ...)
            current_rec = {}
            for line in lines:
                if ":" in line:
                    k, v = line.split(":", 1)
                    k_clean = k.strip().lower().replace(" ", "_")
                    v_clean = v.strip()
                    current_rec[k_clean] = v_clean
                    if "applicant_name" in current_rec and "service_type" in current_rec and "department" in current_rec:
                        records.append(current_rec)
                        current_rec = {}
            if current_rec and "applicant_name" in current_rec:
                records.append(current_rec)

    if not records:
        # Fallback single record from PDF text
        records = [{
            "applicant_name": file.filename.replace(".pdf", "").replace("_", " ").title(),
            "service_type": "Document Verification",
            "department": "General Administration",
            "purpose": text[:200]
        }]

    total = len(records)
    imported = 0
    failed = 0
    errors = []

    for i, record in enumerate(records, start=1):
        app, error = _import_single_record(record, db)
        if error:
            failed += 1
            errors.append(f"PDF Record {i}: {error}")
        else:
            imported += 1

    db.commit()

    return {
        "total_rows": total,
        "imported": imported,
        "failed": failed,
        "errors": errors[:20],
    }


@router.post("/ai-analyze")
async def ai_analyze_document(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """
    Analyze uploaded document with Gemini 2.5 Flash AI.
    Categorizes domain (Transport, Scholarship, Real Estate, etc.),
    validates relevance (rejects non-government files), calculates SLA Risk Prediction,
    and automatically creates an Application in the DB if valid.
    """
    from app.services.gemini_service import analyze_with_gemini

    content = await file.read()
    analysis = analyze_with_gemini(file.filename, content)

    if not analysis.get("is_valid", True):
        return {
            "success": False,
            "is_valid": False,
            "filename": file.filename,
            "rejection_reason": analysis.get("rejection_reason", "This document cannot be uploaded. It is an irrelevant file and not a valid government application."),
            "analysis": analysis
        }

    # Save to database if valid
    now = datetime.utcnow()

    record = {
        "applicant_name": analysis.get("applicant_name") or file.filename.split(".")[0].replace("_", " ").title(),
        "service_type": analysis.get("service_type") or "Document Verification",
        "department": analysis.get("department") or "General Administration",
        "district": analysis.get("district") or "North District",
        "purpose": analysis.get("purpose") or f"Imported document: {file.filename}",
        "sla_days": int(analysis.get("sla_days", 15)),
        "submission_date": now.strftime("%Y-%m-%d"),
    }

    application, error = _import_single_record(record, db)
    if error:
        raise HTTPException(status_code=400, detail=f"Failed to create application: {error}")

    # Override risk level & score from Gemini analysis
    if analysis.get("risk_level"):
        application.risk_level = analysis["risk_level"].upper()
    if analysis.get("risk_score"):
        application.risk_score = float(analysis["risk_score"])
    if analysis.get("predicted_delay") is not None:
        application.predicted_delay = bool(analysis["predicted_delay"])
    if analysis.get("predicted_delay_days"):
        application.predicted_delay_days = int(analysis["predicted_delay_days"])
    if analysis.get("risk_factors"):
        application.risk_factors = json.dumps(analysis["risk_factors"])

    # Add document to application
    doc = Document(
        application_id=application.id,
        name=file.filename,
        verified=True,
        size=f"{round(len(content) / 1024, 1)} KB"
    )
    db.add(doc)
    db.commit()
    db.refresh(application)

    from app.api.applications import _to_response
    app_dict = _to_response(application)

    return {
        "success": True,
        "is_valid": True,
        "filename": file.filename,
        "category": analysis.get("category", "General Administration"),
        "analysis": analysis,
        "application": app_dict
    }


